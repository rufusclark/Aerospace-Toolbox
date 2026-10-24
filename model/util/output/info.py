from __future__ import annotations
from .base import count, Path, OutputType

from typing import Any, Concatenate, ParamSpec, TypeVar, Callable, Literal, Generic
from functools import wraps

import pandas as pd
from openpyxl.styles import Side, Border, Font, Alignment


__all__ = [
    "scientific_table",
    "Table"
]

P = ParamSpec("P")
T = TypeVar("T", bound="OutputType")

#
#
# Table
#
#

SaveFormats = Literal[".csv", ".xlsx", ".tex"]


class Table(Generic[T, P]):
    def __init__(
        self,
        func: Callable[Concatenate[T, P], pd.DataFrame] | None = None,
        path: Path | None = None,
        obj: T | None = None,
        superposed_tables: list[Table[T, P]] | None = None
    ) -> None:
        self.func = func
        self.path = path
        self.obj = obj

        self.superposed_tables = superposed_tables if superposed_tables else []

        self.args: P.args = ()  # type: ignore
        self.kwargs: P.kwargs = {}  # type: ignore

        if func is not None:
            wraps(func)(self)

    def __get__(
            self,
            obj: T | None,
            owner: type[T] | None = None
    ) -> Table[T, P]:
        if obj is None:
            return self

        return Table(
            func=self.func,
            path=self.path,
            obj=obj,
            superposed_tables=self.superposed_tables
        )

    def __repr__(self) -> str:
        if self.path:
            return f"<Table {self.path.title}>"
        elif self.func:
            return f"<Table {self.func.__name__}>"
        else:
            raise ValueError(
                "neither self.path or self.func are defined."
            )

    def __call__(
        self,
        *args: P.args,
        save: bool = True,  # type: ignore
        show: bool = True,  # type: ignore
        transpose: bool = False,  # type: ignore
        suppress_prints: bool = False,  # type: ignore
        extensions: tuple[SaveFormats, ...] = (  # type: ignore
            ".csv", ".xlsx", ".tex"),
        ** kwargs: P.kwargs
    ) -> Table[T, P]:

        if self.obj is None or self.func is None:
            raise TypeError(
                "Table must be accessed through an instance before calling it"
            )
        obj = self.obj

        self.args = args
        self.kwargs = kwargs

        if self.path is None:
            self.path = Path(
                instance_name=obj.model.name,
                method_name=self.func.__name__,
                category_name=obj.model.category,
                output_type=obj.output_type,
                kwargs=kwargs
            )

        if save:
            self.save(extensions=extensions,
                      suppress_prints=suppress_prints, transpose=transpose)

        if show:
            self.show(suppress_prints=suppress_prints, transpose=transpose)

        count.tables += 1

        return self

    def __add__(self, other: Table | None) -> Table:
        if other is None:
            return self

        if not isinstance(other, Table):
            raise TypeError(
                f"{type(other)} is not supported for this operation"
            )

        assert self.path is not None
        assert other.path is not None

        return Table(
            func=self.func,
            path=self.path + other.path,
            obj=self.obj,
            superposed_tables=[other, *self.superposed_tables,
                               *other.superposed_tables],
        )

    def __radd__(self, other: Table | None):
        return self.__add__(other)

    def df(self) -> pd.DataFrame:
        assert self.func
        assert self.obj
        assert self.path

        df = self.func(self.obj, *self.args, **self.kwargs)  # type: ignore

        if self.superposed_tables:
            if not "Name" in df.columns:
                df.insert(0, "Name", self.path.instance_name)

        for other in self.superposed_tables:
            assert other.func
            assert other.obj
            assert other.path

            df1 = other.func(
                other.obj, *other.args, **other.kwargs
            )  # type: ignore
            if not "Name" in df1.columns:
                df1.insert(0, "Name", other.path.instance_name)

            df = pd.concat([df, df1], ignore_index=True)

        return df

    def __transpose(self, df: pd.DataFrame) -> pd.DataFrame:
        out_df = df.T.reset_index()
        out_df.columns = out_df.iloc[0]
        out_df = out_df[1:].reset_index(drop=True)

        return out_df

    def __sort(self, df: pd.DataFrame, column: str, ascending: bool = True) -> pd.DataFrame:
        out_df = df.sort_values(
            by=column,
            ascending=ascending
        )
        return out_df

    def save(
            self,
            transpose: bool = False,
            sort_column: str | None = None,
            sort_ascending: bool = True,
            extensions: tuple[SaveFormats, ...] = (".csv", ".xlsx", ".tex"), suppress_prints: bool = False
    ):
        if ".csv" in extensions:
            self.save_csv(suppress_prints=True, transpose=transpose,
                          sort_column=sort_column, sort_ascending=sort_ascending)
        if ".xlsx" in extensions:
            self.save_excel(suppress_prints=True, transpose=transpose,
                            sort_column=sort_column, sort_ascending=sort_ascending)
        if ".tex" in extensions:
            self.save_latex(suppress_prints=True, transpose=transpose,
                            sort_column=sort_column, sort_ascending=sort_ascending)

        if not suppress_prints:
            assert self.path
            print(f"Saved tables to {self.path.to_stem}")

    def save_csv(
            self,
            transpose: bool = False,
            sort_column: str | None = None,
            sort_ascending: bool = True,
            suppress_prints: bool = False
    ):
        df = self.df()

        if sort_column is not None:
            df = self.__sort(df, sort_column, sort_ascending)
        if transpose:
            df = self.__transpose(df)

        assert self.path is not None

        self.path.parent.mkdir(parents=True, exist_ok=True)
        path = self.path.to_path(".csv")
        df.to_csv(path, index=False)

        if not suppress_prints:
            print(f"Saved table to {path}")

    def save_latex(
        self,
        transpose: bool = False,
        sort_column: str | None = None,
        sort_ascending: bool = True,
        suppress_prints: bool = False
    ):
        df = self.df()

        if sort_column is not None:
            df = self.__sort(df, sort_column, sort_ascending)
        if transpose:
            df = self.__transpose(df)

        assert self.path is not None

        self.path.parent.mkdir(parents=True, exist_ok=True)
        path = self.path.to_path(".tex")
        df.to_latex(path, index=False)

        if not suppress_prints:
            print(f"Saved table to {path}")

    def save_excel(
        self,
        transpose: bool = False,
        sort_column: str | None = None,
        sort_ascending: bool = True,
        suppress_prints: bool = False
    ):
        df = self.df()

        if sort_column is not None:
            df = self.__sort(df, sort_column, sort_ascending)
        if transpose:
            df = self.__transpose(df)

        assert self.path is not None

        self.path.parent.mkdir(parents=True, exist_ok=True)
        path = self.path.to_path(".xlsx")

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Data")

            ws = writer.sheets["Data"]

            thin = Side(style="thin")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            # fmt numbers and add borders
            for row in ws.iter_rows():
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="top")

                    value = cell.value
                    if not isinstance(value, (int, float)):
                        continue

                    if round(value, 2) != 0:
                        cell.number_format = "0.00"

                    if value == 0:
                        cell.number_format = "0.00"

                    if round(value, 2) == 0:
                        exponent = int(f"{value:.0e}".split("e")[1])
                        coeff = value / (10 ** exponent)

                        superscript = str(exponent).translate(
                            str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")
                        )

                        cell.number_format = "@"
                        cell.value = f"{coeff:.2f} × 10{superscript}"
                        cell.alignment = Alignment(horizontal="right")

            # set font
            for column in ws.columns:
                for cell in column:
                    cell.font = Font(name="Maven Pro")

            # bold first row
            for cell in ws["1"]:
                cell.font = Font(name="Maven Pro", bold=True)

            # auto-size columns
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter

                for cell in column:
                    if cell.value is not None:
                        max_length = max(max_length, len(str(cell.value)))

                ws.column_dimensions[column_letter].width = max_length

        if not suppress_prints:
            print(f"Saved table to {path}")

    def show(
        self,
        transpose: bool = False,
        sort_column: str | None = None,
        sort_ascending: bool = True,
        suppress_prints: bool = False
    ):
        df = self.df()

        if sort_column is not None:
            df = self.__sort(df, sort_column, sort_ascending)
        if transpose:
            df = self.__transpose(df)

        if not suppress_prints and self.path is not None:
            print(self.path.title)

        print(df.to_string(
            index=False,
            float_format="{:.3f}".format
        ))

# TODO: Delete old path implementation
# TODO: Delete old decorator implementation


def scientific_table() -> Callable[
    [Callable[Concatenate[T, P], pd.DataFrame]],
    Table[T, P]
]:
    def decorator(
        func: Callable[Concatenate[T, P], pd.DataFrame]
    ) -> Table[T, P]:

        return Table(
            func=func
        )

    return decorator
